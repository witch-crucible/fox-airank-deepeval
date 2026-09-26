import json
import os
import tempfile
import threading
import unittest
import warnings
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.parse import quote
from urllib.request import Request, urlopen

from model_dashboard.server import (
    ARENA_DATASET_URL,
    ARTIFICIAL_ANALYSIS_URL,
    LLM_STATS_INDEX_URL,
    DashboardError,
    SameOriginRedirectHandler,
    DashboardStore,
    calculate_scores,
    collect_watch_snapshot,
    create_server,
    fetch_json,
    main,
    normalize_agent_usage_entry,
    normalize_model,
    parse_agent_usage_csv,
    resolve_static_path,
    url_origin,
    normalize_arena_webdev_rows,
    normalize_artificial_analysis_html,
    normalize_external_rows,
    normalize_llm_stats_indexes,
)
from model_dashboard.leaderboards import (
    build_agent_overview,
    build_local_scores,
    build_model_overview,
    normalize_third_party_weights,
)
from model_dashboard.sources import normalize_arena_price_catalog, normalize_artificial_analysis_models_html


class ModelDashboardTests(unittest.TestCase):
    def test_server_facade_reexports_split_module_api(self):
        from model_dashboard import domain, sources, storage
        from model_dashboard import server as server_module

        self.assertIs(server_module.DashboardError, domain.DashboardError)
        self.assertIs(server_module.calculate_scores, domain.calculate_scores)
        self.assertIs(server_module.normalize_model, domain.normalize_model)
        self.assertIs(server_module.fetch_json, sources.fetch_json)
        self.assertIs(
            server_module.normalize_arena_webdev_rows,
            sources.normalize_arena_webdev_rows,
        )
        self.assertIs(server_module.DashboardStore, storage.DashboardStore)

    @staticmethod
    def _artificial_analysis_document(count=3, score_offset=0):
        rows = []
        for index in range(1, count + 1):
            score = (70 - index + score_offset) / 100
            rows.append(
                {
                    "id": f"aa-{index}",
                    "display": {
                        "agent": f"Agent {index}",
                        "model": f"Model {index}",
                        "creator": {"agent": "Test Lab", "model": "Test Lab"},
                    },
                    "displayLabel": f"Agent {index} - Model {index}",
                    "indexScore": score,
                    "evals": [
                        {"datasetIndexName": "deep-swe", "mean": {"reward": score - 0.01}},
                        {"datasetIndexName": "terminal-bench-v2.1", "mean": {"reward": score + 0.02}},
                        {"datasetIndexName": "swe-atlas-qna", "mean": {"reward": score - 0.01}},
                    ],
                    "mean": {"costUsd": index / 10, "agentWallTimeSec": index * 100},
                }
            )
        payload = json.dumps([1, "prefix" + "".join(json.dumps(row) for row in rows)])
        return (
            "<html><body>Artificial Analysis Coding Agent Index v1.3"
            f"<script>self.__next_f.push({payload})</script></body></html>"
        )

    @staticmethod
    def _arena_document(count=31, score_offset=0):
        return {
            "rows": [
                {
                    "row": {
                        "model_name": f"Arena Model {rank}",
                        "organization": "Test Lab",
                        "license": "Test License",
                        "rating": 1701 - rank + score_offset,
                        "rank": rank,
                        "category": "overall",
                        "leaderboard_publish_date": "2026-08-12",
                    }
                }
                for rank in range(1, count + 1)
            ]
        }

    @staticmethod
    def _llm_stats_document(count=3, score_offset=0):
        def rows(category_offset):
            return [
                {
                    "model_id": f"llm-{rank}",
                    "model_name": f"LLM {rank}",
                    "organization_id": "test-lab",
                    "organization_name": "Test Lab",
                    "conservative": 60 - rank + category_offset + score_offset,
                    "rank": rank,
                    "rank_delta_14d": 1 - rank,
                    "games_played": 10 + rank,
                }
                for rank in range(1, count + 1)
            ]

        return {
            "general": {"models": rows(0)},
            "reasoning": {"models": rows(-1)},
            "code": {"models": rows(-2)},
            "agents": {"models": rows(-3)},
        }

    def test_calculate_scores_preserves_missing_values_and_uses_case_weights(self):
        scores = calculate_scores(
            {
                "skill_call": 10,
                "code_review": 8,
                "logic_analysis": 9,
                "function_fix": 9,
                "model_test_correction": 100,
                "model_test_generation": 100,
                "model_test_logic": 87.14,
                "arena_webdev": 1565,
            }
        )
        self.assertEqual(scores["manual_total"], 36)
        self.assertEqual(scores["model_test_total"], 93.14)
        self.assertEqual(scores["composite_total"], 1694.14)

        missing = calculate_scores({})
        self.assertIsNone(missing["manual_total"])
        self.assertIsNone(missing["model_test_total"])
        self.assertIsNone(missing["composite_total"])

    def test_external_rows_use_dot_paths_and_reject_unknown_targets(self):
        models = normalize_external_rows(
            {"data": {"models": [{"identity": {"agent": "Codex", "model": "GPT-X"}, "score": 96.5}]}},
            "data.models",
            {"tool": "identity.agent", "model": "identity.model", "scores.model_test_total": "score"},
            "测试接口",
            "https://example.com/models",
        )
        self.assertEqual(models[0]["tool"], "Codex")
        self.assertEqual(models[0]["scores"]["model_test_total"], 96.5)
        self.assertEqual(models[0]["source"]["type"], "third_party")

        with self.assertRaisesRegex(DashboardError, "不支持的目标字段"):
            normalize_external_rows([], "", {"tool": "tool", "model": "model", "unsafe": "x"}, "source", "https://example.com")

    def test_arena_rows_normalize_full_board_and_allow_explicit_limit(self):
        models = normalize_arena_webdev_rows(self._arena_document())
        self.assertEqual(len(models), 31)
        self.assertEqual(models[0]["model"], "Arena Model 1")
        self.assertEqual(models[0]["scores"]["arena_webdev"], 1700)
        self.assertEqual(models[0]["source"]["type"], "arena_webdev")
        self.assertEqual(models[-1]["source"]["rank"], 31)

        limited = normalize_arena_webdev_rows(self._arena_document(), limit=30)
        self.assertEqual(len(limited), 30)

        with self.assertRaisesRegex(DashboardError, "不足 30 条"):
            normalize_arena_webdev_rows(self._arena_document(count=29), limit=30)

    def test_arena_official_catalog_normalizes_model_price_aliases(self):
        prices = normalize_arena_price_catalog([{
            "name": "Alpha Display",
            "model_api_name": "alpha-api",
            "input_token_price": "1.25",
            "output_token_price": "5",
            "price_source": "https://example.com/pricing",
        }])
        self.assertEqual(prices["alpha-api"]["input"], 1.25)
        self.assertEqual(prices["alpha display"]["output"], 5)

    def test_artificial_analysis_html_normalizes_score_and_rank(self):
        models = normalize_artificial_analysis_html(self._artificial_analysis_document())
        self.assertEqual(len(models), 3)
        self.assertEqual(models[0]["tool"], "Agent 1")
        self.assertEqual(models[0]["model"], "Model 1")
        self.assertEqual(models[0]["scores"]["artificial_analysis_index"], 69)
        self.assertEqual(models[0]["scores"]["aa_deep_swe"], 68)
        self.assertEqual(models[0]["source"]["rank"], 1)
        self.assertEqual(models[0]["source"]["index_version"], "v1.3")

        with self.assertRaisesRegex(DashboardError, "缺少榜单数据"):
            normalize_artificial_analysis_html("<html></html>")

    def test_artificial_analysis_rows_select_all_beyond_top_30(self):
        models = normalize_artificial_analysis_html(self._artificial_analysis_document(count=35))
        self.assertEqual(len(models), 35)
        self.assertEqual(models[0]["tool"], "Agent 1")
        self.assertEqual(models[0]["model"], "Model 1")
        self.assertEqual(models[-1]["tool"], "Agent 35")
        self.assertEqual(models[-1]["model"], "Model 35")
        self.assertEqual(models[-1]["source"]["rank"], 35)
        for model in models:
            self.assertEqual(model["source"]["type"], "artificial_analysis_agent")
            scores = model["scores"]
            self.assertIsInstance(scores["artificial_analysis_index"], (int, float))
            self.assertIsInstance(scores["aa_deep_swe"], (int, float))
            self.assertIsInstance(scores["aa_terminal_bench_v2"], (int, float))
            self.assertIsInstance(scores["aa_swe_atlas_qna"], (int, float))

        truncated = normalize_artificial_analysis_html(self._artificial_analysis_document(count=35), limit=30)
        self.assertEqual(len(truncated), 30)
        self.assertEqual(truncated[-1]["source"]["rank"], 30)

    def test_artificial_analysis_accepts_legacy_terminal_bench_v2_key(self):
        document = self._artificial_analysis_document(count=1).replace(
            "terminal-bench-v2.1",
            "terminal-bench-v2",
        )
        models = normalize_artificial_analysis_html(document)
        self.assertEqual(len(models), 1)
        self.assertIsInstance(models[0]["scores"]["aa_terminal_bench_v2"], (int, float))

    def test_llm_stats_indexes_normalize_official_rank_and_scores(self):
        models = normalize_llm_stats_indexes(self._llm_stats_document())
        self.assertEqual(len(models), 3)
        self.assertEqual(models[0]["tool"], "Test Lab")
        self.assertEqual(models[0]["model"], "LLM 1")
        self.assertEqual(models[0]["scores"]["llm_stats_score"], 59)
        self.assertEqual(models[0]["scores"]["llm_stats_reasoning"], 58)
        self.assertEqual(models[0]["scores"]["llm_stats_code"], 57)
        self.assertEqual(models[0]["scores"]["llm_stats_agents"], 56)
        self.assertEqual(models[0]["source"]["rank"], 1)
        self.assertEqual(models[0]["source"]["source_id"], "llm-1")

        with self.assertRaisesRegex(DashboardError, "缺少 general 总榜"):
            normalize_llm_stats_indexes({})

    def test_llm_stats_indexes_keep_full_board_and_allow_explicit_limit(self):
        models = normalize_llm_stats_indexes(self._llm_stats_document(count=35))
        self.assertEqual(len(models), 35)
        self.assertEqual(models[0]["model"], "LLM 1")
        self.assertEqual(models[-1]["model"], "LLM 35")
        self.assertEqual(models[-1]["source"]["rank"], 35)
        self.assertEqual(models[-1]["source"]["source_id"], "llm-35")
        limited = normalize_llm_stats_indexes(self._llm_stats_document(count=35), limit=30)
        self.assertEqual(len(limited), 30)

    def test_artificial_analysis_models_uses_full_manifests_and_keeps_variants(self):
        references = (
            '"manifest":{"path":"/data/canonical.txt","key":"' + "01" * 32 + '"}'
            '"manifest":{"path":"/data/hosts.txt","key":"' + "02" * 32 + '"}'
        )
        payload = json.dumps([1, references])
        document = (
            "<html>Artificial Analysis Intelligence Index v4.3.2"
            f"<script>self.__next_f.push({payload})</script></html>"
        )
        canonical = {
            "models": [{
                "id": "model-1",
                "slug": "alpha-1",
                "name": "Alpha 1",
                "intelligenceIndex": 72.5,
                "terminalBench40": 0.44,
                "timescaleData": {"medianOutputSpeed": 123.4},
                "creator": {"name": "Test Lab"},
            }]
        }
        hosts = [{
            "id": "host-1",
            "slug": "lab_alpha-1",
            "modelId": "model-1",
            "modelSlug": "alpha-1",
            "host": {"name": "Provider"},
            "intelligenceIndexCostPerTask": {"cost": {"total": 0.42}},
            "timescaleData": {"medianOutputSpeed": 111.0},
        }]
        with patch("model_dashboard.sources.decrypt_artificial_analysis_manifest", side_effect=[canonical, hosts]):
            models = normalize_artificial_analysis_models_html(document, lambda _url: b"encrypted")
        self.assertEqual(len(models), 1)
        self.assertEqual(models[0]["source"]["type"], "artificial_analysis_model")
        self.assertEqual(models[0]["scores"]["aa_model_intelligence"], 72.5)
        self.assertEqual(models[0]["scores"]["aa_model_terminal_bench_v4"], 44)
        self.assertEqual(models[0]["source"]["slug"], "alpha-1")
        self.assertEqual(models[0]["source"]["variants"][0]["cost_usd_per_task"], 0.42)

    def test_model_overview_applies_requested_weights_and_local_eligibility(self):
        def model(source_type, name, score_field, score, rank):
            return normalize_model(
                {"tool": "Lab", "model": name, "reasoning_effort": "max", "scores": {score_field: score}},
                source={"type": source_type, "source_id": f"{source_type}-{rank}", "rank": rank},
            )

        models = [
            model("arena_webdev", "Alpha (max)", "arena_webdev", 1800, 1),
            model("arena_webdev", "Beta (max)", "arena_webdev", 1600, 2),
            model("artificial_analysis_model", "Alpha", "aa_model_intelligence", 80, 1),
            model("artificial_analysis_model", "Beta", "aa_model_intelligence", 60, 2),
            model("llm_stats", "Alpha", "llm_stats_score", 50, 1),
            model("llm_stats", "Beta", "llm_stats_score", 40, 2),
        ]
        metrics = {
            name: {"score": 0.9}
            for name in ("Task Correctness", "Robustness, Safety and Regression", "Delivery Evidence")
        }
        local_records = [
            {"agent": "Agent", "model": "Alpha", "reasoning_effort": "max", "case_id": case_id, "metrics": metrics}
            for case_id in ("case-a", "case-b")
            for _ in range(3)
        ]
        overview = build_model_overview(models, local_records, {"case-a", "case-b"})
        alpha = next(row for row in overview["models"] if row["model_key"] == "alpha")
        self.assertEqual(alpha["third_party_score"], 100)
        self.assertEqual(alpha["local_configurations"][0]["local_score"], 90)
        self.assertEqual(alpha["local_configurations"][0]["battle_score"], 97)
        self.assertEqual(overview["weights"], {"arena_webdev": 0.35, "artificial_analysis_model": 0.5, "llm_stats": 0.15})

        incomplete = build_local_scores(local_records[:2], {"case-a", "case-b"})[0]
        self.assertFalse(incomplete["eligible"])
        self.assertEqual(incomplete["missing_cases"], ["case-b"])

    def test_model_overview_applies_configured_third_party_weights(self):
        def model(source_type, name, score_field, score, rank):
            return normalize_model(
                {"tool": "Lab", "model": name, "scores": {score_field: score}},
                source={"type": source_type, "source_id": f"{source_type}-{rank}", "rank": rank},
            )

        models = [
            model("arena_webdev", "Alpha", "arena_webdev", 100, 1),
            model("arena_webdev", "Beta", "arena_webdev", 0, 2),
            model("artificial_analysis_model", "Alpha", "aa_model_intelligence", 0, 2),
            model("artificial_analysis_model", "Beta", "aa_model_intelligence", 100, 1),
            model("llm_stats", "Alpha", "llm_stats_score", 0, 2),
            model("llm_stats", "Beta", "llm_stats_score", 100, 1),
        ]
        weights = {"artificial_analysis_model": 0.1, "arena_webdev": 0.8, "llm_stats": 0.1}
        overview = build_model_overview(models, [], set(), weights=weights)

        self.assertEqual(overview["weights"], weights)
        self.assertEqual([row["model"] for row in overview["models"]], ["Alpha", "Beta"])
        self.assertEqual([row["third_party_score"] for row in overview["models"]], [80, 20])

    def test_third_party_weights_require_all_sources_and_total_100_percent(self):
        self.assertEqual(normalize_third_party_weights(), {
            "artificial_analysis_model": 0.5,
            "arena_webdev": 0.35,
            "llm_stats": 0.15,
        })
        with self.assertRaisesRegex(DashboardError, "必须包含"):
            normalize_third_party_weights({"arena_webdev": 1})
        with self.assertRaisesRegex(DashboardError, "合计必须为 100%"):
            normalize_third_party_weights({
                "artificial_analysis_model": 0.5,
                "arena_webdev": 0.3,
                "llm_stats": 0.1,
            })

    def test_zero_weight_source_is_not_required_for_model_ranking(self):
        arena = normalize_model(
            {"tool": "Arena", "model": "Alpha", "scores": {"arena_webdev": 1800}},
            source={"type": "arena_webdev", "source_id": "alpha", "rank": 1},
        )
        overview = build_model_overview(
            [arena],
            [],
            set(),
            weights={"artificial_analysis_model": 0, "arena_webdev": 1, "llm_stats": 0},
        )

        self.assertEqual(overview["models"][0]["third_party_score"], 100)
        self.assertEqual(overview["models"][0]["third_party_rank"], 1)

    def test_model_overview_matches_aa_release_without_merging_efforts(self):
        def model(source_type, name, score_field, score, rank, **source):
            return normalize_model(
                {"tool": "Anthropic", "model": name, "scores": {score_field: score}},
                source={"type": source_type, "source_id": f"{source_type}-{rank}-{name}", "rank": rank, **source},
            )

        models = [
            model("arena_webdev", "claude-fable-5.1-max", "arena_webdev", 1700, 1),
            model("arena_webdev", "claude-opus-5-max", "arena_webdev", 1650, 2),
            model(
                "artificial_analysis_model",
                "Claude Fable 5.1 (Adaptive Reasoning, Max Effort, Default Fallback)",
                "aa_model_intelligence",
                53,
                1,
                creator_name="Anthropic",
                release={"slug": "claude-fable-5-1", "name": "Claude Fable 5.1"},
            ),
            model(
                "artificial_analysis_model",
                "Claude Fable 5.1 (Adaptive Reasoning, High Effort, Default Fallback)",
                "aa_model_intelligence",
                51,
                3,
                creator_name="Anthropic",
                release={"slug": "claude-fable-5-1", "name": "Claude Fable 5.1"},
            ),
            model(
                "artificial_analysis_model",
                "Claude Opus 5 (Adaptive Reasoning, Max Effort)",
                "aa_model_intelligence",
                50,
                2,
                creator_name="Anthropic",
                release={"slug": "claude-opus-5", "name": "Claude Opus 5"},
            ),
            model("llm_stats", "Claude Fable 5.1", "llm_stats_score", 56, 1),
            model("llm_stats", "Claude Opus 5", "llm_stats_score", 54, 2),
        ]

        overview = build_model_overview(models, [], set())
        rows = {(row["model_key"], row["reasoning_effort"]): row for row in overview["models"]}

        self.assertEqual(set(rows), {
            ("claude fable 5 1", "max"), ("claude fable 5 1", "high"), ("claude fable 5 1", "unknown"),
            ("claude opus 5", "max"), ("claude opus 5", "unknown"),
        })
        fable_max = rows[("claude fable 5 1", "max")]
        fable_high = rows[("claude fable 5 1", "high")]
        self.assertEqual(fable_max["model"], "Claude Fable 5.1")
        self.assertEqual(fable_max["sources"]["artificial_analysis_model"]["raw_score"], 53)
        self.assertEqual(fable_high["sources"]["artificial_analysis_model"]["raw_score"], 51)
        self.assertEqual(len(fable_max["sources"]["artificial_analysis_model"]["variants"]), 1)
        self.assertNotIn("llm_stats", fable_max["sources"])
        self.assertNotIn("arena_webdev", fable_high["sources"])
        self.assertTrue(all(row["third_party_score"] is None for row in rows.values()))

    def test_agent_overview_matches_model_name_and_quantifies_shared_terminal_bench(self):
        baseline = normalize_model(
            {"tool": "Anthropic", "model": "Claude Opus 5 (max)", "reasoning_effort": "max", "scores": {"aa_model_terminal_bench_v4": 40}},
            source={"type": "artificial_analysis_model", "source_id": "model", "creator_name": "Anthropic", "variants": []},
        )
        agent = normalize_model(
            {"tool": "Claude Code", "model": "Opus 5 (max)", "scores": {"aa_terminal_bench_v4": 55}},
            source={"type": "artificial_analysis_agent", "source_id": "agent", "rank": 1, "creator": {"model": "Anthropic"}},
        )
        overview = build_agent_overview([baseline, agent])
        self.assertEqual(overview["agents"][0]["baseline"]["model"], "Claude Opus 5 (max)")
        self.assertEqual(overview["agents"][0]["terminal_bench_uplift"], 15)

    def test_store_keeps_leaderboard_snapshots_and_compares_previous(self):
        with tempfile.TemporaryDirectory() as directory:
            data_path = Path(directory) / "dashboard.json"
            data_path.write_text(json.dumps({"meta": {}, "models": []}), encoding="utf-8")
            store = DashboardStore(data_path=data_path)
            first = normalize_arena_webdev_rows(self._arena_document(count=2))
            second = normalize_arena_webdev_rows(self._arena_document(count=2, score_offset=10))
            store.sync_arena_webdev(first)
            store.sync_arena_webdev(second)
            snapshots = store.leaderboard_snapshots("arena_webdev")
            self.assertEqual(len(snapshots), 2)
            comparison = store.leaderboard_comparison("arena_webdev")
            self.assertEqual(comparison["current"]["snapshot_id"], snapshots[0]["snapshot_id"])
            self.assertEqual(comparison["rows"][0]["metric_deltas"]["arena_webdev"], 10)
            alias = store.save_model_alias("Opus 5", "Claude Opus 5")
            self.assertEqual(alias["canonical_key"], "claude opus 5")
            self.assertEqual(store.read()["model_aliases"]["opus 5"], "claude opus 5")
            weights = store.save_leaderboard_weights({
                "artificial_analysis_model": 0.2,
                "arena_webdev": 0.7,
                "llm_stats": 0.1,
            })
            self.assertEqual(weights["arena_webdev"], 0.7)
            self.assertEqual(store.leaderboard_weights(), weights)
            self.assertEqual(store.read()["meta"]["leaderboard_weights"], weights)

    def test_store_seeds_existing_leaderboard_as_legacy_history_baseline(self):
        with tempfile.TemporaryDirectory() as directory:
            data_path = Path(directory) / "dashboard.json"
            existing = normalize_arena_webdev_rows(self._arena_document(count=2))
            data_path.write_text(json.dumps({"meta": {}, "models": existing}), encoding="utf-8")
            store = DashboardStore(data_path=data_path)
            snapshots = store.leaderboard_snapshots("arena_webdev")
            self.assertEqual(len(snapshots), 1)
            self.assertEqual(snapshots[0]["row_count"], 2)

    def test_http_page_add_model_and_import_third_party_data(self):
        third_party_document = {
            "payload": [
                {
                    "agent": "OpenCode",
                    "name": "Remote Model",
                    "scores": {"model_test": 91, "arena": 1500},
                }
            ]
        }
        fetched = []

        def fake_fetcher(url, auth):
            fetched.append((url, auth))
            return third_party_document

        with tempfile.TemporaryDirectory() as directory:
            data_path = Path(directory) / "dashboard.json"
            server = create_server("127.0.0.1", 0, data_path=data_path, json_fetcher=fake_fetcher)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                with urlopen(f"{base_url}/", timeout=3) as response:
                    html = response.read().decode("utf-8")
                    self.assertIn("模型能力台", html)
                    self.assertIn("/static/brand/mark.png", html)
                    self.assertIn("/static/brand/favicon.ico", html)
                    self.assertIn("增加模型", html)
                    self.assertIn("录入使用人数", html)
                    self.assertIn("导入使用人数 CSV", html)
                    self.assertIn("更多数据操作", html)
                    self.assertIn("section-nav", html)
                    self.assertIn("panel-kicker", html)
                    self.assertIn("Agent 使用人数", html)
                    self.assertIn("/api/agent-usage", html)
                    self.assertIn("/api/agent-usage/import", html)
                    self.assertIn("openAgentUsageImportDialog", html)
                    self.assertIn("看板模式", html)
                    self.assertIn("表格模式", html)
                    self.assertIn("按归档状态筛选", html)
                    self.assertIn('aria-sort', html)
                    self.assertIn('按${column.label}排序', html)
                    self.assertIn('每页 50 条', html)
                    self.assertIn('本地更新', html)
                    self.assertIn('role="status"', html)
                    self.assertNotIn('<main id="app" aria-live=', html)
                    self.assertIn("同步三方数据", html)
                    self.assertIn("sync-all-leaderboards", html)
                    self.assertIn("runAllLeaderboardSyncs", html)
                    self.assertIn("按来源单独同步", html)
                    self.assertIn("同步 Arena 全榜", html)
                    self.assertIn("同步 AA Agent 全榜", html)
                    self.assertIn("同步 AA Model 全榜", html)
                    self.assertNotIn("同步 AA 前 30", html)
                    self.assertIn("同步 LLM Stats 全榜", html)
                    self.assertNotIn(">同步榜单</summary>", html)
                    self.assertIn("能力排名", html)
                    self.assertIn("Agent 排名", html)
                    self.assertIn("Model 排名", html)
                    self.assertIn("function renderRankings(models)", html)
                    self.assertIn("function bestAgentEntries(models)", html)
                    self.assertIn("function scoredModels(models)", html)
                    self.assertIn("function appendRankingRows(chart, ranked, titleFn, subtitleFn)", html)
                    self.assertIn("ranking-split", html)
                    self.assertIn("按 AI 编程工具聚合，取该指标最佳成绩", html)
                    self.assertIn("按模型配置逐条比较当前指标", html)
                    self.assertIn("配对模型：", html)
                    self.assertIn("配对 Agent：", html)
                    self.assertIn("空白不按 0 分处理", html)
                    self.assertIn('data-view="table"', html)
                    self.assertIn('data-view="recommendations"', html)
                    self.assertIn('data-view="recommendation-log"', html)
                    self.assertIn('data-view="agent-reference"', html)
                    self.assertIn("function renderModelLeaderboardPage()", html)
                    self.assertIn("function renderAgentLeaderboardPage()", html)
                    self.assertIn("function sourceDetailHref(model)", html)
                    self.assertIn('parsed.hash = "artificial-analysis-coding-agent-index"', html)
                    self.assertIn('source.slug || source.release?.slug', html)
                    self.assertIn('parsed.pathname = `/models/${encodeURIComponent(slug)}`', html)
                    self.assertIn("function leaderboardMetricLink(value, href, title)", html)
                    self.assertIn("sourceDetailHref(agent)", html)
                    self.assertIn("sourceDetailHref(baseline)", html)
                    self.assertIn('`TB4 ${number(baseline.scores?.aa_model_terminal_bench_v4)}`', html)
                    self.assertNotIn('el("strong", "", baseline.model)', html)
                    self.assertIn("function leaderboardPodiumClass(rank)", html)
                    self.assertIn("leaderboard-podium-1", html)
                    self.assertIn("leaderboardPodiumClass(displayRank)", html)
                    self.assertIn("纯 Artificial Analysis", html)
                    self.assertIn("AA 官方名次", html)
                    self.assertIn("leaderboardPodiumClass(source.rank)", html)
                    self.assertIn("function leaderboardModelHref(source)", html)
                    self.assertIn('target.pathname = `/models/${encodeURIComponent(slug)}`', html)
                    self.assertIn('target.pathname = `/models/${encodeURIComponent(sourceId)}`', html)
                    self.assertIn('target.searchParams.set("q", model)', html)
                    self.assertIn('heading.rel = "noopener noreferrer"', html)
                    self.assertIn('data-view="settings"', html)
                    self.assertIn("function orderedModelLeaderboardSources(weights = {})", html)
                    self.assertIn("function renderSettingsPage()", html)
                    self.assertIn('api("/api/leaderboards/weights"', html)
                    self.assertIn("Model 三方榜单权重", html)
                    self.assertIn("三个来源均可设为 0%，但合计必须为 100%", html)
                    self.assertIn('heading: "Arena", defaultWeight: .35', html)
                    self.assertIn('view: "recommendations"', html)
                    self.assertIn('function renderRecommendationsPage()', html)
                    self.assertIn('function renderRecommendationLogPage()', html)
                    self.assertIn('function renderLegionArtwork(src, alt, summaryText)', html)
                    self.assertIn('el("details", "legion-artwork")', html)
                    self.assertIn('展开军团海报', html)
                    self.assertIn('展开技能冷却现场', html)
                    self.assertIn('/static/art/vibe-coding-legion.png', html)
                    self.assertIn('/static/art/vibe-coding-legion-log.png', html)
                    self.assertIn('function recommendationModel(entry)', html)
                    self.assertIn('function recommendationLocalRecord(entry)', html)
                    self.assertIn('function recommendationPurpose(entry, group)', html)
                    self.assertIn('function recommendationPurposeTypes()', html)
                    self.assertIn('function recommendationAgentName(entry, planTitle)', html)
                    self.assertIn('function recommendationEntries()', html)
                    self.assertIn('function recommendationAgentGroups(entries)', html)
                    self.assertIn('function recommendationPurposeGroups(entries)', html)
                    self.assertIn('function pricingEntryForTool(tool)', html)
                    self.assertIn('function recommendationPricingLink(tool)', html)
                    self.assertIn('link.href = "#pricing"', html)
                    self.assertIn('state.pricingTargetTool = pricing.tool', html)
                    self.assertIn('recommendation-pricing-link', html)
                    self.assertIn('price-card.is-targeted', html)
                    self.assertIn('已定位 ${targetTool} 的订阅费用', html)
                    self.assertIn('purposeRole: "primary"', html)
                    self.assertIn('purposeRole: "secondary"', html)
                    self.assertIn('is-secondary-purpose', html)
                    self.assertIn('function recommendationModelLogo(entry)', html)
                    self.assertIn('key: "openai"', html)
                    self.assertIn('key: "claude"', html)
                    self.assertIn('key: "deepseek"', html)
                    self.assertIn('recommendation-model-logo', html)
                    self.assertIn('function recommendationPurposeType(entry, group)', html)
                    self.assertIn('function recommendationPrimaryPurposeTypes(entry, group)', html)
                    self.assertIn('function recommendationSecondaryPurposeTypes(entry)', html)
                    self.assertIn('function recommendationSecondaryPurposeType(entry)', html)
                    self.assertIn('function recommendationSelectedPurposeTypes(entry, group)', html)
                    self.assertIn('function renderRecommendationCapability(entry, metrics)', html)
                    self.assertIn('用途类型配置 *', html)
                    self.assertIn('["Ask", "Plan", "Build", "Review", "Ship"]', html)
                    self.assertIn('主用途（可多选）', html)
                    self.assertIn('副用途（可多选）', html)
                    self.assertIn('function renderRecommendationPurposeOptions(field, selected = [], inputName = "primary_purpose_types")', html)
                    self.assertIn('container.dataset.recommendationGroup', html)
                    self.assertIn('const selectedSecondaryPurposeTypes = recommendationSecondaryPurposeTypes(entry)', html)
                    self.assertIn('function filterRecommendationEditorRows()', html)
                    self.assertIn('function syncRecommendationPurposeRoles(row, changedInput)', html)
                    self.assertIn('function requestCloseRecommendationsDialog()', html)
                    self.assertIn('recommendation-editor-filter', html)
                    self.assertIn('recommendation-editor-summary', html)
                    self.assertIn('recommendation-editor-row-actions', html)
                    self.assertIn('正在保存…', html)
                    self.assertIn('有尚未保存的 Vibe Coding Legion 修改', html)
                    self.assertIn("record.primary_purpose_types", html)
                    self.assertIn("record.secondary_purpose_types", html)
                    self.assertIn('function recommendationCorePurposeTypes(entry, group)', html)
                    self.assertIn('function toggleRecommendationCore(group, index, purposeType, nextValue, control)', html)
                    self.assertIn('recommendation-core-toggle', html)
                    self.assertIn('recommendation-item.is-core', html)
                    self.assertIn('isCorePurpose ? "★" : "☆"', html)
                    self.assertIn('coreToggle.setAttribute("aria-label"', html)
                    self.assertIn('已标记为核心选择', html)
                    self.assertIn('"核心用途（按用途 + 条目，可多选）"', html)
                    self.assertIn('tag.title = "主用途"', html)
                    self.assertIn('tag.title = "副用途"', html)
                    self.assertIn('用途说明（可选）', html)
                    self.assertIn('recommendation-purpose', html)
                    self.assertIn('暂无匹配的能力数据，空白不按 0 分处理', html)
                    self.assertIn('card.dataset.purposeGroup', html)
                    self.assertIn('agentGroup.dataset.agentGroup', html)
                    self.assertIn('按主用途分组', html)
                    self.assertIn('每个用途内部按 Agent 聚合', html)
                    self.assertIn('不按 Agent 能力归类', html)
                    self.assertIn("function sourceHref(model)", html)
                    self.assertIn("function sourceBadge(model)", html)
                    self.assertIn("打开来源：", html)
                    self.assertIn("function pricingPlanGroups(item)", html)
                    self.assertIn("Agent Plan", html)
                    self.assertIn("Coding Plan", html)
                    self.assertIn("agent_plans", html)
                    self.assertIn("coding_plans", html)
                    self.assertIn("ide_plans", html)
                    self.assertIn("code_plans", html)
                    self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")

                with urlopen(f"{base_url}/static/brand/mark.png", timeout=3) as response:
                    self.assertEqual(response.headers["Content-Type"], "image/png")
                    self.assertGreater(len(response.read()), 0)

                for asset in ("vibe-coding-legion.png", "vibe-coding-legion-log.png"):
                    with self.subTest(asset=asset):
                        with urlopen(f"{base_url}/static/art/{asset}", timeout=3) as response:
                            self.assertEqual(response.status, 200)
                            self.assertEqual(response.headers["Content-Type"], "image/png")
                            self.assertGreater(len(response.read()), 0)

                with self.assertRaises(HTTPError) as blocked:
                    urlopen(f"{base_url}/static/../seed_data.json", timeout=3)
                self.assertEqual(blocked.exception.code, 404)

                created = self._post_json(
                    f"{base_url}/api/models",
                    {
                        "id": "client-controlled-id",
                        "tool": "Codex",
                        "model": "Local Model",
                        "scores": {"skill_call": 8},
                        "source": {"type": "arena_webdev", "name": "伪造来源"},
                        "archived": True,
                    },
                )
                self.assertEqual(created["model"]["scores"]["manual_total"], 8)
                self.assertNotEqual(created["model"]["id"], "client-controlled-id")
                self.assertEqual(created["model"]["source"], {"type": "manual", "name": "手工录入"})
                self.assertFalse(created["model"]["archived"])

                imported = self._post_json(
                    f"{base_url}/api/import",
                    {
                        "name": "远端榜单",
                        "url": "https://example.com/models.json",
                        "array_path": "payload",
                        "mapping": {
                            "tool": "agent",
                            "model": "name",
                            "scores.model_test_total": "scores.model_test",
                            "scores.arena_webdev": "scores.arena",
                        },
                    },
                )
                self.assertEqual(imported, {"created": 1, "updated": 0, "skipped": 0, "received": 1})
                self.assertEqual(fetched, [("https://example.com/models.json", None)])

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                self.assertEqual(len(data["models"]), 22)
                self.assertEqual(data["models"][-1]["source"]["name"], "远端榜单")
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_authenticated_json_source_requires_https_before_network_access(self):
        with patch.dict(os.environ, {"MODEL_API_TOKEN": "secret"}, clear=False):
            with (
                patch("model_dashboard.sources.socket.getaddrinfo") as mocked_resolver,
                patch("model_dashboard.sources.urlopen") as mocked_urlopen,
            ):
                with self.assertRaisesRegex(DashboardError, "必须使用 HTTPS"):
                    fetch_json("http://example.com/models.json", {"env": "MODEL_API_TOKEN"})
                mocked_resolver.assert_not_called()
                mocked_urlopen.assert_not_called()

    def test_authenticated_redirect_rejects_cross_origin_before_copying_headers(self):
        request = Request(
            "https://source.example/models.json",
            headers={"Authorization": "Bearer secret"},
        )
        handler = SameOriginRedirectHandler()
        with self.assertRaisesRegex(DashboardError, "不能重定向到其他站点"):
            handler.redirect_request(
                request,
                None,
                302,
                "Found",
                {},
                "https://other.example/models.json",
            )

    def test_authenticated_redirect_treats_default_https_port_as_same_origin(self):
        request = Request(
            "https://source.example/models.json",
            headers={"Authorization": "Bearer secret"},
        )
        redirected = SameOriginRedirectHandler().redirect_request(
            request,
            None,
            302,
            "Found",
            {},
            "https://source.example:443/next.json",
        )
        self.assertIsNotNone(redirected)
        self.assertEqual(redirected.get_header("Authorization"), "Bearer secret")

    def test_normalize_model_rejects_all_non_object_sources(self):
        for source in ([], False, 0, "", "bad"):
            with self.subTest(source=source):
                with self.assertRaisesRegex(DashboardError, "source 必须是对象"):
                    normalize_model({"tool": "Codex", "model": "Model", "source": source})

    def test_url_origin_rejects_invalid_ports_with_dashboard_error(self):
        for url in ("https://example.com:bad/models", "https://example.com:70000/models"):
            with self.subTest(url=url):
                with self.assertRaisesRegex(DashboardError, "URL 端口无效"):
                    url_origin(url)

    def test_resolve_static_path_allows_brand_assets_and_blocks_traversal(self):
        for asset in ("mark.png", "apple-touch-icon.png", "favicon.ico"):
            with self.subTest(asset=asset):
                brand_asset = resolve_static_path(f"/static/brand/{asset}")
                self.assertIsNotNone(brand_asset)
                assert brand_asset is not None
                self.assertEqual(brand_asset.name, asset)
        for asset in ("vibe-coding-legion.png", "vibe-coding-legion-log.png"):
            with self.subTest(asset=asset):
                image = resolve_static_path(f"/static/art/{asset}")
                self.assertIsNotNone(image)
                assert image is not None
                self.assertEqual(image.name, asset)
        self.assertIsNone(resolve_static_path("/static/../seed_data.json"))
        self.assertIsNone(resolve_static_path("/static/index.html"))
        self.assertIsNone(resolve_static_path("/static/brand/"))
        self.assertIsNone(resolve_static_path("/static/brand/missing.svg"))

    def test_archive_rejects_duplicate_legacy_ids_without_modifying_data(self):
        with tempfile.TemporaryDirectory() as directory:
            data_path = Path(directory) / "dashboard.json"
            data = {
                "meta": {},
                "models": [
                    {"id": "duplicate", "archived": False, "updated_at": "first"},
                    {"id": "duplicate", "archived": False, "updated_at": "second"},
                ],
            }
            data_path.write_text(json.dumps(data), encoding="utf-8")
            store = DashboardStore(data_path=data_path)

            with self.assertRaisesRegex(DashboardError, "包含重复的模型 ID"):
                store.set_model_archived("duplicate", True)

            self.assertEqual(json.loads(data_path.read_text(encoding="utf-8")), data)

    def test_http_archives_and_restores_model_without_deleting_it(self):
        with tempfile.TemporaryDirectory() as directory:
            server = create_server("127.0.0.1", 0, data_path=Path(directory) / "dashboard.json")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                created = self._post_json(
                    f"{base_url}/api/models",
                    {"tool": "Codex", "model": "Archive Model"},
                )["model"]
                archived = self._post_json(
                    f"{base_url}/api/models/{created['id']}/archive",
                    {"archived": True},
                )["model"]
                self.assertTrue(archived["archived"])
                self.assertTrue(archived["archived_at"])

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                stored = next(model for model in data["models"] if model["id"] == created["id"])
                self.assertTrue(stored["archived"])

                restored = self._post_json(
                    f"{base_url}/api/models/{created['id']}/archive",
                    {"archived": False},
                )["model"]
                self.assertFalse(restored["archived"])
                self.assertEqual(restored["archived_at"], "")
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_http_archive_decodes_percent_encoded_model_id(self):
        """前端用 encodeURIComponent 拼 ID，归档接口必须解码后再落库。"""
        with tempfile.TemporaryDirectory() as directory:
            server = create_server("127.0.0.1", 0, data_path=Path(directory) / "dashboard.json")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                created = self._post_json(f"{base_url}/api/models", {"tool": "Codex", "model": "Opus 5.5 (max)"})["model"]
                encoded = quote(created["id"], safe="")
                archived = self._post_json(f"{base_url}/api/models/{encoded}/archive", {"archived": True})["model"]
                self.assertTrue(archived["archived"])
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_third_party_overwrite_preserves_archived_state(self):
        with tempfile.TemporaryDirectory() as directory:
            data_path = Path(directory) / "dashboard.json"
            server = create_server("127.0.0.1", 0, data_path=data_path)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                created = self._post_json(
                    f"{base_url}/api/models",
                    {"tool": "Codex", "model": "Archive Model", "scores": {"skill_call": 5}},
                )["model"]
                self._post_json(
                    f"{base_url}/api/models/{created['id']}/archive",
                    {"archived": True},
                )
                replacement = {
                    **created,
                    "scores": {"skill_call": 9},
                    "source": {"type": "third_party", "name": "测试接口"},
                }
                result = server.store.import_models([replacement], overwrite=True)  # type: ignore[attr-defined]
                self.assertEqual(result, {"created": 0, "updated": 1, "skipped": 0})

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                stored = next(model for model in data["models"] if model["id"] == created["id"])
                self.assertTrue(stored["archived"])
                self.assertTrue(stored["archived_at"])
                self.assertEqual(stored["scores"]["skill_call"], 9)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_http_syncs_full_arena_without_overwriting_other_models(self):
        documents = [self._arena_document(), self._arena_document(score_offset=5)]
        fetched = []

        def fake_fetcher(url, auth):
            fetched.append((url, auth))
            return documents.pop(0)

        with tempfile.TemporaryDirectory() as directory:
            server = create_server("127.0.0.1", 0, data_path=Path(directory) / "dashboard.json", json_fetcher=fake_fetcher)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                first = self._post_json(f"{base_url}/api/import/arena-webdev", {})
                self.assertEqual(first, {"created": 31, "updated": 0, "removed": 0, "received": 31})
                second = self._post_json(f"{base_url}/api/import/arena-webdev", {})
                self.assertEqual(second, {"created": 0, "updated": 31, "removed": 0, "received": 31})
                self.assertEqual(fetched, [(ARENA_DATASET_URL, None), (ARENA_DATASET_URL, None)])

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                arena_models = [model for model in data["models"] if model["source"]["type"] == "arena_webdev"]
                self.assertEqual(len(arena_models), 31)
                self.assertEqual(arena_models[0]["scores"]["arena_webdev"], 1705)
                self.assertEqual(len(data["models"]), 51)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_arena_sync_retains_archived_model_after_it_leaves_full_board(self):
        first_document = self._arena_document()
        second_document = self._arena_document()
        second_document["rows"] = [
            row
            for row in second_document["rows"]
            if row["row"]["model_name"] != "Arena Model 30"
        ]
        documents = [first_document, second_document]

        def fake_fetcher(url, auth):
            return documents.pop(0)

        with tempfile.TemporaryDirectory() as directory:
            server = create_server(
                "127.0.0.1",
                0,
                data_path=Path(directory) / "dashboard.json",
                json_fetcher=fake_fetcher,
            )
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                self._post_json(f"{base_url}/api/import/arena-webdev", {})
                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                archived_model = next(model for model in data["models"] if model["model"] == "Arena Model 30")
                self._post_json(
                    f"{base_url}/api/models/{archived_model['id']}/archive",
                    {"archived": True},
                )

                result = self._post_json(f"{base_url}/api/import/arena-webdev", {})
                self.assertEqual(result, {"created": 0, "updated": 30, "removed": 0, "received": 30})
                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                retained = next(model for model in data["models"] if model["model"] == "Arena Model 30")
                self.assertTrue(retained["archived"])
                self.assertEqual(len([model for model in data["models"] if model["source"]["type"] == "arena_webdev"]), 31)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_http_syncs_artificial_analysis_ranking_without_overwriting_other_models(self):
        documents = [
            self._artificial_analysis_document(),
            self._artificial_analysis_document(score_offset=5),
            self._artificial_analysis_document(count=35),
        ]
        fetched = []

        def fake_text_fetcher(url):
            fetched.append(url)
            return documents.pop(0)

        with tempfile.TemporaryDirectory() as directory:
            server = create_server(
                "127.0.0.1",
                0,
                data_path=Path(directory) / "dashboard.json",
                text_fetcher=fake_text_fetcher,
            )
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                first = self._post_json(f"{base_url}/api/import/artificial-analysis", {})
                self.assertEqual(first, {"created": 3, "updated": 0, "removed": 0, "received": 3})
                second = self._post_json(f"{base_url}/api/import/artificial-analysis", {})
                self.assertEqual(second, {"created": 0, "updated": 3, "removed": 0, "received": 3})

                select_all = self._post_json(f"{base_url}/api/import/artificial-analysis", {})
                self.assertEqual(select_all, {"created": 32, "updated": 3, "removed": 0, "received": 35})
                self.assertEqual(
                    fetched,
                    [ARTIFICIAL_ANALYSIS_URL, ARTIFICIAL_ANALYSIS_URL, ARTIFICIAL_ANALYSIS_URL],
                )

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                aa_models = [model for model in data["models"] if model["source"]["type"] == "artificial_analysis_agent"]
                non_aa = [model for model in data["models"] if model["source"]["type"] != "artificial_analysis_agent"]
                self.assertEqual(len(aa_models), 35)
                self.assertGreater(len(aa_models), 30)
                self.assertEqual(aa_models[0]["scores"]["artificial_analysis_index"], 69)
                self.assertEqual(aa_models[-1]["source"]["rank"], 35)
                self.assertEqual(len(non_aa), 20)
                self.assertEqual(len(data["models"]), 55)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_http_syncs_llm_stats_ranking_and_preserves_previous_data_on_invalid_response(self):
        documents = [self._llm_stats_document(), self._llm_stats_document(score_offset=5), {}]
        fetched = []

        def fake_fetcher(url, auth):
            fetched.append((url, auth))
            return documents.pop(0)

        with tempfile.TemporaryDirectory() as directory:
            server = create_server(
                "127.0.0.1",
                0,
                data_path=Path(directory) / "dashboard.json",
                json_fetcher=fake_fetcher,
            )
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                first = self._post_json(f"{base_url}/api/import/llm-stats", {})
                self.assertEqual(first, {"created": 3, "updated": 0, "removed": 0, "received": 3})
                second = self._post_json(f"{base_url}/api/import/llm-stats", {})
                self.assertEqual(second, {"created": 0, "updated": 3, "removed": 0, "received": 3})
                with self.assertRaises(HTTPError) as raised:
                    self._post_json(f"{base_url}/api/import/llm-stats", {})
                self.assertEqual(raised.exception.code, 400)
                raised.exception.close()
                self.assertEqual(fetched, [(LLM_STATS_INDEX_URL, None)] * 3)

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                llm_stats_models = [model for model in data["models"] if model["source"]["type"] == "llm_stats"]
                self.assertEqual(len(llm_stats_models), 3)
                self.assertEqual(llm_stats_models[0]["scores"]["llm_stats_score"], 64)
                self.assertEqual(len(data["models"]), 23)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_http_exposes_model_agent_history_and_alias_apis(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            runs_root = root / "runs"
            cases_root = root / "cases"
            runs_root.mkdir()
            cases_root.mkdir()
            server = create_server(
                "127.0.0.1",
                0,
                data_path=root / "dashboard.json",
                runs_root=runs_root,
                cases_root=cases_root,
            )
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                with urlopen(f"{base_url}/api/leaderboards/models", timeout=3) as response:
                    models = json.load(response)
                self.assertEqual(models["weights"]["arena_webdev"], 0.35)
                self.assertEqual(models["local_share"], 0.3)

                with urlopen(f"{base_url}/api/leaderboards/weights", timeout=3) as response:
                    self.assertEqual(json.load(response)["weights"], models["weights"])
                configured = self._post_json(
                    f"{base_url}/api/leaderboards/weights",
                    {"weights": {
                        "artificial_analysis_model": 0.2,
                        "arena_webdev": 0.7,
                        "llm_stats": 0.1,
                    }},
                )
                self.assertEqual(configured["weights"]["arena_webdev"], 0.7)
                with urlopen(f"{base_url}/api/leaderboards/models", timeout=3) as response:
                    self.assertEqual(json.load(response)["weights"], configured["weights"])
                with self.assertRaises(HTTPError) as raised:
                    self._post_json(
                        f"{base_url}/api/leaderboards/weights",
                        {"weights": {
                            "artificial_analysis_model": 0.2,
                            "arena_webdev": 0.2,
                            "llm_stats": 0.2,
                        }},
                    )
                self.assertEqual(raised.exception.code, 400)
                raised.exception.close()

                with urlopen(f"{base_url}/api/leaderboards/agents", timeout=3) as response:
                    agents = json.load(response)
                self.assertGreater(agents["count"], 0)
                self.assertTrue(all(row["supplemental"] for row in agents["agents"]))
                self.assertTrue(all(row["legion"]["matched"] for row in agents["agents"]))
                self.assertTrue(all(not row["agent"]["scores"] for row in agents["agents"]))

                with urlopen(f"{base_url}/api/leaderboards/arena_webdev", timeout=3) as response:
                    history = json.load(response)
                self.assertEqual(history["rows"], [])

                saved = self._post_json(
                    f"{base_url}/api/leaderboards/model-aliases",
                    {"alias": "Opus 5", "canonical": "Claude Opus 5"},
                )
                self.assertEqual(saved["alias"]["canonical_key"], "claude opus 5")
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_http_rejects_invalid_model(self):
        with tempfile.TemporaryDirectory() as directory:
            server = create_server("127.0.0.1", 0, data_path=Path(directory) / "dashboard.json")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                with self.assertRaises(HTTPError) as raised:
                    self._post_json(
                        f"http://127.0.0.1:{server.server_port}/api/models",
                        {"tool": "", "model": "Missing tool"},
                    )
                self.assertEqual(raised.exception.code, 400)
                error = json.loads(raised.exception.read().decode("utf-8"))
                self.assertEqual(error["error"], "AI 编程工具不能为空")
                raised.exception.close()
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_normalize_agent_usage_entry_requires_non_negative_count(self):
        entry = normalize_agent_usage_entry(
            {"tool": "  Claude Code  ", "user_count": 1200, "notes": "内部统计"}
        )
        self.assertEqual(entry["tool"], "Claude Code")
        self.assertEqual(entry["user_count"], 1200)
        self.assertEqual(entry["notes"], "内部统计")
        self.assertTrue(entry["updated_at"])

        with self.assertRaisesRegex(DashboardError, "AI 编程工具不能为空"):
            normalize_agent_usage_entry({"tool": " ", "user_count": 1})
        with self.assertRaisesRegex(DashboardError, "使用人数不能为空"):
            normalize_agent_usage_entry({"tool": "Codex", "user_count": None})
        with self.assertRaisesRegex(DashboardError, "使用人数不能为负数"):
            normalize_agent_usage_entry({"tool": "Codex", "user_count": -1})
        with self.assertRaisesRegex(DashboardError, "user_count 必须是数字"):
            normalize_agent_usage_entry({"tool": "Codex", "user_count": "很多"})

    def test_store_upserts_agent_usage_by_case_insensitive_tool(self):
        with tempfile.TemporaryDirectory() as directory:
            data_path = Path(directory) / "dashboard.json"
            data_path.write_text(json.dumps({"meta": {}, "models": []}), encoding="utf-8")
            store = DashboardStore(data_path=data_path)

            created, is_new = store.upsert_agent_usage({"tool": "Claude Code", "user_count": 10})
            self.assertTrue(is_new)
            self.assertEqual(created["user_count"], 10)

            updated, is_new = store.upsert_agent_usage(
                {"tool": "claude code", "user_count": 25, "notes": "季度统计"}
            )
            self.assertFalse(is_new)
            self.assertEqual(updated["tool"], "claude code")
            self.assertEqual(updated["user_count"], 25)

            data = json.loads(data_path.read_text(encoding="utf-8"))
            self.assertEqual(len(data["agent_usage"]), 1)
            self.assertEqual(data["agent_usage"][0]["user_count"], 25)
            self.assertEqual(data["agent_usage"][0]["notes"], "季度统计")
            self.assertTrue(data["meta"]["updated_at"])

    def test_http_upserts_agent_usage_and_exposes_it_on_models_payload(self):
        with tempfile.TemporaryDirectory() as directory:
            server = create_server("127.0.0.1", 0, data_path=Path(directory) / "dashboard.json")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                created = self._post_json(
                    f"{base_url}/api/agent-usage",
                    {"tool": "Codex", "user_count": 88, "notes": "试点"},
                )
                self.assertTrue(created["created"])
                self.assertEqual(created["entry"]["tool"], "Codex")
                self.assertEqual(created["entry"]["user_count"], 88)

                updated = self._post_json(
                    f"{base_url}/api/agent-usage",
                    {"tool": "codex", "user_count": 120},
                )
                self.assertFalse(updated["created"])
                self.assertEqual(updated["entry"]["user_count"], 120)

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                self.assertEqual(len(data["agent_usage"]), 1)
                self.assertEqual(data["agent_usage"][0]["tool"], "codex")
                self.assertEqual(data["agent_usage"][0]["user_count"], 120)

                with self.assertRaises(HTTPError) as raised:
                    self._post_json(f"{base_url}/api/agent-usage", {"tool": "Codex", "user_count": -3})
                self.assertEqual(raised.exception.code, 400)
                raised.exception.close()
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_parse_agent_usage_csv_supports_chinese_headers_and_last_row_wins(self):
        entries = parse_agent_usage_csv(
            "\ufeff工具,使用人数,备注\n"
            'Claude Code,"1,000",初版\n'
            "Codex,800,\n"
            "claude code,1500,覆盖\n"
            ",,\n"
        )
        self.assertEqual(len(entries), 2)
        by_tool = {entry["tool"].casefold(): entry for entry in entries}
        self.assertEqual(by_tool["claude code"]["user_count"], 1500)
        self.assertEqual(by_tool["claude code"]["notes"], "覆盖")
        self.assertEqual(by_tool["codex"]["user_count"], 800)

        with self.assertRaisesRegex(DashboardError, "表头需包含"):
            parse_agent_usage_csv("name,score\nCodex,1\n")
        with self.assertRaisesRegex(DashboardError, "第 2 行"):
            parse_agent_usage_csv("tool,user_count\nCodex,-3\n")
        with self.assertRaisesRegex(DashboardError, "没有有效的使用人数数据行"):
            parse_agent_usage_csv("tool,user_count\n,,\n")

    def test_http_imports_agent_usage_csv_and_overwrites_same_tool(self):
        with tempfile.TemporaryDirectory() as directory:
            server = create_server("127.0.0.1", 0, data_path=Path(directory) / "dashboard.json")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base_url = f"http://127.0.0.1:{server.server_port}"
            try:
                self._post_json(
                    f"{base_url}/api/agent-usage",
                    {"tool": "Codex", "user_count": 10, "notes": "旧值"},
                )
                result = self._post_json(
                    f"{base_url}/api/agent-usage/import",
                    {
                        "csv": (
                            "tool,user_count,notes\n"
                            "Codex,120,CSV 覆盖\n"
                            "Claude Code,90,\n"
                        )
                    },
                )
                self.assertEqual(result["received"], 2)
                self.assertEqual(result["created"], 1)
                self.assertEqual(result["updated"], 1)

                with urlopen(f"{base_url}/api/models", timeout=3) as response:
                    data = json.load(response)
                usage = {entry["tool"]: entry for entry in data["agent_usage"]}
                self.assertEqual(usage["Codex"]["user_count"], 120)
                self.assertEqual(usage["Codex"]["notes"], "CSV 覆盖")
                self.assertEqual(usage["Claude Code"]["user_count"], 90)

                with self.assertRaises(HTTPError) as raised:
                    self._post_json(f"{base_url}/api/agent-usage/import", {"csv": "tool,user_count\n"})
                self.assertEqual(raised.exception.code, 400)
                raised.exception.close()
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=3)

    def test_seed_data_matches_excel_source(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("未安装可选的 openpyxl，跳过 Excel 源文件对账")
        workbook_path = Path("/Users/ben/Downloads/模型能力测试-20260623-展示优化版.xlsx")
        if not workbook_path.exists():
            self.skipTest("原始 Excel 不在当前机器上")
        seed_path = Path(__file__).resolve().parent.parent / "model_dashboard" / "seed_data.json"
        seed = json.loads(seed_path.read_text(encoding="utf-8"))
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            sheet = load_workbook(workbook_path, data_only=True)["测试结果"]
        self.assertEqual(len(seed["models"]), sheet.max_row - 1)
        for row, model in zip(sheet.iter_rows(min_row=2, values_only=True), seed["models"]):
            self.assertEqual([model["tool"], model["model"], model["reasoning_effort"]], [row[0] or "", row[1] or "", row[2] or ""])
            source_scores = {
                "manual_total": row[7] if row[7] != "" else None,
                "model_test_total": row[11] if row[11] != "" else None,
                "arena_webdev": row[12] if row[12] != "" else None,
            }
            for field, value in source_scores.items():
                self.assertEqual(model["scores"][field], value)
            has_component = any(value is not None for value in source_scores.values())
            expected_composite = row[13] if has_component else None
            self.assertEqual(model["scores"]["composite_total"], expected_composite)

    def test_seed_pricing_uses_agent_and_coding_plan_groups(self):
        seed_path = Path(__file__).resolve().parent.parent / "model_dashboard" / "seed_data.json"
        pricing = json.loads(seed_path.read_text(encoding="utf-8"))["pricing"]

        self.assertTrue(pricing)
        for entry in pricing:
            self.assertIn("agent_plans", entry)
            self.assertIn("coding_plans", entry)
            self.assertNotIn("ide_plans", entry)
            self.assertNotIn("code_plans", entry)
            self.assertNotIn("plans", entry)

    def test_collect_watch_snapshot_tracks_python_and_ignores_local_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            watched = root / "server.py"
            ignored = root / "data.local.json"
            cached = root / "__pycache__" / "server.cpython-312.pyc"
            static_html = root / "static" / "index.html"
            cached.parent.mkdir()
            static_html.parent.mkdir()
            watched.write_text("print('ok')\n", encoding="utf-8")
            ignored.write_text("{}\n", encoding="utf-8")
            cached.write_bytes(b"cache")
            static_html.write_text("<html></html>\n", encoding="utf-8")

            snapshot = collect_watch_snapshot(root)
            self.assertIn(str(watched.resolve()), snapshot)
            self.assertNotIn(str(ignored.resolve()), snapshot)
            self.assertNotIn(str(cached.resolve()), snapshot)
            self.assertNotIn(str(static_html.resolve()), snapshot)

            watched.write_text("print('changed')\n", encoding="utf-8")
            os.utime(watched, ns=(snapshot[str(watched.resolve())] + 1_000_000, snapshot[str(watched.resolve())] + 1_000_000))
            updated = collect_watch_snapshot(root)
            self.assertNotEqual(snapshot[str(watched.resolve())], updated[str(watched.resolve())])

    def test_main_reload_flag_delegates_to_run_with_reload(self):
        with patch("model_dashboard.server.run_with_reload") as mocked_reload:
            main(["--reload", "--host", "0.0.0.0", "--port", "9000", "--data", "/tmp/dashboard.json"])
        mocked_reload.assert_called_once_with(
            ["--host", "0.0.0.0", "--port", "9000", "--data", "/tmp/dashboard.json"]
        )

    @staticmethod
    def _post_json(url, payload):
        request = Request(
            url,
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urlopen(request, timeout=3) as response:
            return json.load(response)


if __name__ == "__main__":
    unittest.main()
